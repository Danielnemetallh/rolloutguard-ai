from pathlib import Path

from openpyxl import load_workbook

from rolloutguard_api.services.export import sanitize_cell


def test_sanitize_formula_injection() -> None:
    assert sanitize_cell("=cmd|'/c calc'!A0") == "'=cmd|'/c calc'!A0"
    assert sanitize_cell("+1234") == "'+1234"
    assert sanitize_cell("-1+1") == "'-1+1"
    assert sanitize_cell("@SUM(A1)") == "'@SUM(A1)"
    assert sanitize_cell("DE-NRW-0107") == "DE-NRW-0107"
    assert sanitize_cell(42) == 42


def test_export_analysis_writes_files(tmp_path: Path) -> None:
    from fastapi.testclient import TestClient

    from rolloutguard_api.db.session import SessionLocal
    from rolloutguard_api.main import create_app
    from rolloutguard_api.services.export import export_analysis

    client = TestClient(create_app())
    project_id = client.get("/api/projects").json()[0]["id"]
    analysis_id = client.post(f"/api/projects/{project_id}/analyze-synthetic").json()[
        "analysis_run_id"
    ]

    db = SessionLocal()
    try:
        result = export_analysis(db, analysis_id, output_dir=tmp_path)
    finally:
        db.close()

    xlsx = Path(result["xlsx_path"])
    md = Path(result["markdown_path"])
    assert xlsx.exists()
    assert md.exists()
    wb = load_workbook(xlsx, read_only=True)
    assert "Findings" in wb.sheetnames
    assert result["finding_count"] >= 1
