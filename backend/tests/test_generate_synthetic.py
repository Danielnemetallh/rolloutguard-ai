import json
from pathlib import Path

from openpyxl import load_workbook

from rolloutguard_api.domain.generate_synthetic import generate
from rolloutguard_api.domain.schema import HERO_SITE_ID


def test_generate_synthetic_workbooks(tmp_path: Path) -> None:
    paths = generate(tmp_path, n_sites=50)

    assert paths["contract"].exists()
    assert paths["schedule"].exists()
    assert paths["status"].exists()
    assert paths["manifest"].exists()
    assert paths["sow"].exists()
    assert "DE-NRW-0107" in paths["sow"].name or paths["sow"].suffix == ".pdf"

    manifest = json.loads(paths["manifest"].read_text(encoding="utf-8"))
    assert manifest["site_count"] >= 40
    assert manifest["hero_site_id"] == HERO_SITE_ID
    assert len(manifest["anomalies"]) >= 12

    # Hero site must appear in all three workbooks
    contract = load_workbook(paths["contract"], read_only=True)
    schedule = load_workbook(paths["schedule"], read_only=True)
    status = load_workbook(paths["status"], read_only=True)

    def first_col_ids(wb) -> set[str]:
        ws = wb.active
        rows = ws.iter_rows(min_row=2, max_col=1, values_only=True)
        return {str(r[0]) for r in rows if r and r[0]}

    assert HERO_SITE_ID in first_col_ids(contract)
    assert HERO_SITE_ID in first_col_ids(schedule)
    assert HERO_SITE_ID in first_col_ids(status)

    # Schedule uses drifted header "Forecast Date"
    schedule_headers = [
        c.value for c in next(schedule.active.iter_rows(min_row=1, max_row=1))
    ]
    assert "Forecast Date" in schedule_headers

    # Expected rules include hero findings
    hero = next(a for a in manifest["anomalies"] if a["site_id"] == HERO_SITE_ID)
    assert "SLA-001" in hero["rule_ids"]
    assert "SEQ-002" in hero["rule_ids"]
